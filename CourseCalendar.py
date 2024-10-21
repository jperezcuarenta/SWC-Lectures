from docx import Document
import pandas as pd
from datetime import datetime, timedelta


class CourseCalendar:
    def __init__(self, class_days, start_date, end_date, sections_dict, skip_dates, exam_dates):
        self.class_days = class_days
        self.start_date = datetime.strptime(start_date, "%Y-%m-%d")
        self.end_date = datetime.strptime(end_date, "%Y-%m-%d")
        self.sections_dict = sections_dict
        self.skip_dates = {datetime.strptime(date, "%Y-%m-%d") for date in skip_dates}
        self.exam_dates = {datetime.strptime(date, "%Y-%m-%d"): exam for date, exam in exam_dates.items()}
        self.schedule = []

    def generate_schedule(self):
        num_section_dates = self.num_section_dates()
        distributed_sections = self.distribute_sections(self.sections_dict, num_section_dates)
        current_date = self.start_date
        section_index = 0
        while current_date <= self.end_date:
            if current_date in self.skip_dates:
                self.schedule.append((current_date.strftime("%Y-%m-%d"), "No Lecture"))
            elif current_date in self.exam_dates:
                self.schedule.append((current_date.strftime("%Y-%m-%d"), self.exam_dates[current_date]))
            elif current_date.strftime("%A") in self.class_days:
                if section_index < len(distributed_sections):
                    self.schedule.append((current_date.strftime("%Y-%m-%d"), ', '.join(distributed_sections[section_index])))
                    section_index += 1
            current_date += timedelta(days=1)
        df = pd.DataFrame(self.schedule, columns=["Date", "Agenda"])
        df['Week Number'] = self.get_weekly_idxs()
        df['Day of Week'] = pd.to_datetime(df['Date']).dt.day_name()
        self.write_calendar_to_memory(df)

    def num_section_dates(self):
        current_date = self.start_date
        valid_class_days = 0
        while current_date <= self.end_date:
            if current_date.strftime("%A") in self.class_days and current_date not in self.skip_dates and current_date not in self.exam_dates:
                valid_class_days += 1
            current_date += timedelta(days=1)
        return valid_class_days
    
    def distribute_sections(self, sections_dict, num_section_dates):
        sections_list = list(sections_dict.items())
        total_sections = len(sections_list)
        remainder = total_sections % num_section_dates
        distributed_schedule = [[] for _ in range(num_section_dates)]
        index = 0
        me_thresh = num_section_dates // remainder
        me_cnt = 1
        for day in range(num_section_dates):
            if me_cnt == me_thresh:
                distributed_schedule[day].append(sections_list[index][0] + ': ' + sections_list[index][1])
                if index + 1 == total_sections:
                    break
                distributed_schedule[day].append(sections_list[index + 1][0] + ': ' + sections_list[index + 1][1])
                index += 2
                me_cnt = 1
            else:
                distributed_schedule[day].append(sections_list[index][0] + ': ' + sections_list[index][1])
                index += 1
                me_cnt += 1
        return distributed_schedule

    def get_weekly_idxs(self):
        reference = datetime.strptime(self.schedule[0][0], '%Y-%m-%d').date().isocalendar()[1]
        weekday_cnt = 1
        weekday_id = [weekday_cnt]
        for idx in range(1, len(self.schedule)):
            test_week = datetime.strptime(self.schedule[idx][0], '%Y-%m-%d').date().isocalendar()[1]
            if test_week != reference:
                reference = test_week
                weekday_cnt += 1
            weekday_id.append(weekday_cnt)
        return weekday_id
    
    def write_calendar_to_memory(self, df):
        weekday_to_column_idx = {'Monday': 1, 'Wednesday': 2}
        rows = [[]]
        for _ in range(len(df['Day of Week'].unique())):
            rows.append([])
        for week_num in df['Week Number'].unique():
            df_week = df.loc[df['Week Number'] == week_num]
            rows[0].append(week_num)
            for _, row in df_week.iterrows():
                text = row['Date'] + '\n'
                section = row['Agenda']
                if ',' in section:
                    text += section.split(',')[0] + '\n' + section.split(',')[1][1:]
                else:
                    text += section
                column_idx = weekday_to_column_idx[row['Day of Week']]
                rows[column_idx].append(text)
        data = {'Week Number': rows[0], 'Monday': rows[1], 'Wednesday': rows[2]}
        ms_df = pd.DataFrame(data)
        doc = Document()
        t = doc.add_table(ms_df.shape[0] + 1, ms_df.shape[1])
        for j in range(ms_df.shape[-1]):
            t.cell(0, j).text = ms_df.columns[j]
        for i in range(ms_df.shape[0]):
            for j in range(ms_df.shape[-1]):
                t.cell(i + 1, j).text = str(ms_df.values[i, j])
        doc.save('test.docx')

class_days = ["Monday", "Wednesday"]
start_date = "2024-08-26"
end_date = "2024-12-04"
sections_dict = {'6.2': 'Regions Between Curves',
                 '6.3': 'Volume by Slicing',
                 '6.4': 'Volume by Shells',
                 '6.5': 'Length of Curves',
                 '6.6': 'Surface Area',
                 '6.7': 'Physical Applications',
                 '7.5': 'Trigonometric Functions and Their Inverses', # is actually 1.4
                 '7.6': "L'Hospital's Rule", # is actually 4.7
                 '7.7': 'Hyperbolic Functions', # is 7.3 in the textbook
                 '8.1': 'Basic Approaches',
                 '8.2': 'Integration by Parts',
                 '8.3': 'Trigonometric Integrals',
                 '8.4': 'Trigonometric Substitutions',
                 '8.5': 'Partial Fractions',
                 '8.6': 'Integration Strategies',
                 '8.7': 'Other Methods of Integration',
                 '8.8': 'Numerical Integration',
                 '8.9': 'Improper Integrals',
                 '9.1': 'Basic Ideas',
                 '9.3': 'Separable Differential Equations',
                 '10.1': 'An Overview',
                 '10.2': 'Sequences',
                 '10.3': 'Infinite Series',
                 '10.4': 'Divergence and Integral Tests',
                 '10.5': 'Comparison Tests',
                 '10.6': 'Alternating Series',
                 '10.7': 'The Ratio and Root Tests',
                 '10.8': 'Choosing a Convergence Test',
                 '11.1': 'Approximating Functions with Polynomials',
                 '11.2': 'Properties of Power Series',
                 '11.3': 'Taylor Series',
                 '12.1': 'Parametric Equations',
                 '12.2': 'Polar Coordinates',
                 '12.3': 'Calculus in Polar Coordinates',
                 '12.4': 'Conic Sections'}
skip_dates = ["2024-09-02", "2024-11-11"]
exam_dates = {"2024-09-18": "Exam 01",
              "2024-10-16": "Exam 02",
              "2024-11-13": "Exam 03",}
obj = CourseCalendar(class_days, start_date, end_date, sections_dict, skip_dates, exam_dates)
schedule_df = obj.generate_schedule()